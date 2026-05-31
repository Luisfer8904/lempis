"""
Sub-app privada para Inversiones Guevara Herrera.
Acceso separado usando usuarios IVG desde el mismo login de Lempis.
"""
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from io import BytesIO

from flask import Blueprint, abort, current_app, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import func, inspect, or_
from sqlalchemy.exc import IntegrityError

from models import db
from models.ivg import IVGAgendaItem, IVGCashSummary, IVGClient, IVGPayment, IVGSale, IVGUser


igh_bp = Blueprint("igh", __name__, url_prefix="/igh")

IVG_SALE_CATEGORIES = [
    ("herbicidas", "Herbicidas", "bg-emerald-500"),
    ("concentrados", "Concentrados", "bg-violet-500"),
    ("semillas", "Semillas", "bg-amber-500"),
]
IVG_SALE_CATEGORY_LABELS = {code: label for code, label, _ in IVG_SALE_CATEGORIES}


def _table_exists(model) -> bool:
    return inspect(db.engine).has_table(model.__tablename__)


def _current_igh_user():
    user_id = session.get("igh_user_id")
    if not user_id or not _table_exists(IVGUser):
        return None
    user = db.session.get(IVGUser, user_id)
    if user is None or not user.is_active:
        return None
    return user


def igh_login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = _current_igh_user()
        if user is None:
            session.pop("igh_user_id", None)
            session.pop("igh_user", None)
            session.pop("igh_role", None)
            return redirect(url_for("auth.login"))
        return fn(*args, **kwargs)
    return wrapper


def igh_superadmin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = _current_igh_user()
        if user is None:
            return redirect(url_for("auth.login"))
        if not user.is_superadmin():
            abort(403)
        return fn(*args, **kwargs)
    return wrapper


def igh_admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = _current_igh_user()
        if user is None:
            return redirect(url_for("auth.login"))
        if not (user.is_superadmin() or user.is_admin()):
            abort(403)
        return fn(*args, **kwargs)
    return wrapper


def _can_create_ivg_users(user: IVGUser | None) -> bool:
    return bool(user and (user.is_superadmin() or user.is_admin()))


def _can_edit_ivg_user(actor: IVGUser | None, target: IVGUser | None) -> bool:
    if actor is None or target is None:
        return False
    if actor.is_superadmin():
        return True
    if actor.is_admin():
        if target.id == actor.id:
            return True
        return target.is_cajero()
    return False


def _can_view_ivg_user(actor: IVGUser | None, target: IVGUser | None) -> bool:
    if actor is None or target is None:
        return False
    if actor.is_superadmin():
        return True
    if actor.is_admin():
        return not target.is_superadmin()
    return False


def _can_delete_ivg_user(actor: IVGUser | None, target: IVGUser | None) -> bool:
    if actor is None or target is None:
        return False
    if actor.is_superadmin():
        return target.id != actor.id
    if actor.is_admin():
        return target.is_cajero()
    return False


def _can_access_users_module(user: IVGUser | None) -> bool:
    return bool(user and (user.is_superadmin() or user.is_admin()))


def _can_access_reports_module(user: IVGUser | None) -> bool:
    return bool(user and (user.is_superadmin() or user.is_admin()))


def _can_delete_operational_records(user: IVGUser | None) -> bool:
    return bool(user and (user.is_superadmin() or user.is_admin()))


def _can_edit_ivg_sale(user: IVGUser | None) -> bool:
    return bool(user and (user.is_superadmin() or user.is_admin()))


def _safe_count(model) -> int:
    if not _table_exists(model):
        return 0
    return model.query.count()


def _safe_sum(model, column, *filters):
    if not _table_exists(model):
        return 0
    query = db.session.query(func.coalesce(func.sum(column), 0))
    if filters:
        query = query.filter(*filters)
    return float(query.scalar() or 0)


def _parse_decimal(raw_value: str, label: str) -> Decimal:
    value = (raw_value or "").strip().replace(",", "")
    if not value:
        raise ValueError(f"{label} es obligatorio.")
    try:
        amount = Decimal(value)
    except InvalidOperation as exc:
        raise ValueError(f"{label} no tiene un formato válido.") from exc
    if amount < 0:
        raise ValueError(f"{label} no puede ser negativo.")
    return amount


def _to_decimal(value) -> Decimal:
    if value is None:
        return Decimal("0.00")
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _parse_datetime(raw_value: str, label: str, required: bool = False):
    value = (raw_value or "").strip()
    if not value:
        if required:
            raise ValueError(f"{label} es obligatorio.")
        return None
    for fmt in ("%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    raise ValueError(f"{label} no tiene un formato válido.")


def _format_datetime_local(value):
    if not value:
        return ""
    return value.strftime("%Y-%m-%dT%H:%M")


def _format_date_local(value):
    if not value:
        return ""
    return value.strftime("%Y-%m-%d")


def _format_report_date(value):
    if not value:
        return ""
    return value.strftime("%d/%m/%Y")


def _money_value(value) -> str:
    return f"{_to_decimal(value):.2f}"


def _normalize_ivg_category(value: str | None) -> str:
    category = (value or "herbicidas").strip().lower()
    if category not in IVG_SALE_CATEGORY_LABELS:
        return "herbicidas"
    return category


def _ivg_category_label(value: str | None) -> str:
    category = _normalize_ivg_category(value)
    return IVG_SALE_CATEGORY_LABELS.get(category, category.title())


def _today_start_utc():
    return datetime.combine(datetime.utcnow().date(), datetime.min.time())


def _duplicate_sale_reference(reference_number: str | None, exclude_sale_id: int | None = None):
    reference = (reference_number or "").strip()
    if not reference or not _table_exists(IVGSale):
        return None
    query = IVGSale.query.filter(func.lower(IVGSale.reference_number) == reference.lower())
    if exclude_sale_id is not None:
        query = query.filter(IVGSale.id != exclude_sale_id)
    return query.first()


def _xlsx_response(filename: str, title: str, headers: list[str], rows: list[list[object]]):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Reporte"
    ws.append([title])
    ws.append(["Generado", datetime.utcnow().strftime("%d/%m/%Y")])
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)

    title_fill = PatternFill("solid", fgColor="EEF2FF")
    header_fill = PatternFill("solid", fgColor="E0E7FF")
    ws["A1"].font = Font(bold=True, size=14, color="1E293B")
    ws["A1"].fill = title_fill
    for cell in ws[4]:
        cell.font = Font(bold=True, color="1E293B")
        cell.fill = header_fill
    ws.freeze_panes = "A5"
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 44)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _pdf_response(filename: str, title: str, headers: list[str], rows: list[list[object]]):
    stream = BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=landscape(letter),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Title"]),
        Paragraph(f"Generado: {datetime.utcnow().strftime('%d/%m/%Y')}", styles["Normal"]),
        Spacer(1, 8),
    ]
    table_rows = [headers]
    table_rows.extend([["" if value is None else str(value) for value in row] for row in rows])
    if len(table_rows) == 1:
        table_rows.append(["Sin datos"] + [""] * (len(headers) - 1))

    table = Table(table_rows, colWidths=_pdf_report_widths(len(headers)), repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2FF")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(table)
    doc.build(story)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=filename, mimetype="application/pdf")


def _pdf_report_widths(count: int):
    usable = 258 * mm
    if count <= 4:
        return [usable / count] * count
    first = usable * 0.18
    return [first] + [(usable - first) / (count - 1)] * (count - 1)


def _recalculate_sale_balance(sale: IVGSale) -> None:
    total_paid = sum((payment.amount or Decimal("0.00")) for payment in sale.payments)
    gross_amount = _to_decimal(sale.gross_amount)
    balance = max(gross_amount - _to_decimal(total_paid), Decimal("0.00"))
    sale.balance_due = balance
    if balance == 0:
        sale.status = "pagada"
    elif balance < gross_amount:
        sale.status = "parcial"
    else:
        sale.status = "registrada"


def _month_start():
    now = datetime.utcnow()
    return now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)


def _cashier_missing_summary_dates(days_back: int = 7):
    if not _table_exists(IVGCashSummary):
        return []
    today = datetime.utcnow().date()
    start_date = today - timedelta(days=days_back - 1)
    existing_dates = {
        summary.summary_date.date()
        for summary in IVGCashSummary.query.filter(
            IVGCashSummary.summary_date >= datetime.combine(start_date, datetime.min.time())
        ).all()
    }
    missing_dates = []
    for offset in range(days_back - 1, -1, -1):
        day = today - timedelta(days=offset)
        if day not in existing_dates:
            missing_dates.append(day)
    return missing_dates


def _latest_cash_summary():
    if not _table_exists(IVGCashSummary):
        return None
    return (
        IVGCashSummary.query.order_by(IVGCashSummary.summary_date.desc(), IVGCashSummary.id.desc()).first()
    )


def _client_delete_block_reason(client: IVGClient | None) -> str | None:
    if client is None:
        return "El cliente no existe."

    has_sales = _table_exists(IVGSale) and IVGSale.query.filter_by(client_id=client.id).first() is not None
    has_agenda = _table_exists(IVGAgendaItem) and IVGAgendaItem.query.filter_by(client_id=client.id).first() is not None

    if has_sales or has_agenda:
        return "No puedes eliminar este cliente porque ya tiene facturas, cobros o historial asociado."

    return None


def _calculate_day_close(opening_amount: Decimal, cash_amount: Decimal, withdrawal_amount: Decimal, transfer_amount: Decimal):
    sales_total = cash_amount + transfer_amount
    opening_plus_sales_amount = opening_amount + cash_amount
    expected_close_amount = opening_plus_sales_amount - withdrawal_amount
    return {
        "sales_total": sales_total,
        "opening_plus_sales_amount": opening_plus_sales_amount,
        "expected_close_amount": expected_close_amount,
    }


def _base_context():
    current_igh_user = _current_igh_user()
    month_start = _month_start()
    pending_invoices = []
    recent_collections = []
    recent_cash_summaries = []
    top_clients = []
    if _table_exists(IVGSale):
        pending_invoices = (
            IVGSale.query.filter(
                IVGSale.balance_due > 0,
            )
            .order_by(IVGSale.due_date.asc(), IVGSale.sale_date.desc())
            .limit(8)
            .all()
        )
        top_clients = (
            db.session.query(
                IVGClient,
                func.coalesce(func.sum(IVGSale.balance_due), 0).label("pending_balance"),
            )
            .join(IVGSale, IVGSale.client_id == IVGClient.id)
            .group_by(IVGClient.id)
            .order_by(func.coalesce(func.sum(IVGSale.balance_due), 0).desc())
            .limit(6)
            .all()
        )
    if _table_exists(IVGPayment):
        recent_collections = (
            IVGPayment.query.order_by(IVGPayment.payment_date.desc(), IVGPayment.id.desc())
            .limit(8)
            .all()
        )
    if _table_exists(IVGCashSummary):
        recent_cash_summaries = (
            IVGCashSummary.query.order_by(IVGCashSummary.summary_date.desc(), IVGCashSummary.id.desc())
            .limit(6)
            .all()
        )
    latest_cash_summary = recent_cash_summaries[0] if recent_cash_summaries else None
    latest_variance = _to_decimal(latest_cash_summary.variance_amount) if latest_cash_summary else Decimal("0.00")
    latest_opening_plus_sales = (
        _to_decimal(latest_cash_summary.opening_amount) + _to_decimal(latest_cash_summary.cash_amount)
        if latest_cash_summary else Decimal("0.00")
    )
    category_metrics = [
        {
            "code": code,
            "label": label,
            "color_class": color_class,
            "gross_total": _safe_sum(IVGSale, IVGSale.gross_amount, IVGSale.category == code) if _table_exists(IVGSale) else 0,
            "balance_total": _safe_sum(IVGSale, IVGSale.balance_due, IVGSale.category == code) if _table_exists(IVGSale) else 0,
            "collections_total": _safe_sum(IVGPayment, IVGPayment.amount, IVGPayment.category == code) if _table_exists(IVGPayment) else 0,
        }
        for code, label, color_class in IVG_SALE_CATEGORIES
    ]
    cashier_missing_dates = _cashier_missing_summary_dates()
    pending_agenda_count = (
        IVGAgendaItem.query.filter(IVGAgendaItem.status != "completada").count()
        if _table_exists(IVGAgendaItem)
        else 0
    )
    return {
        "igh_user": current_igh_user.username if current_igh_user else session.get("igh_user"),
        "igh_current_user": current_igh_user,
        "igh_is_superadmin": bool(current_igh_user and current_igh_user.is_superadmin()),
        "igh_is_admin": bool(current_igh_user and current_igh_user.is_admin()),
        "igh_is_cajero": bool(current_igh_user and current_igh_user.is_cajero()),
        "igh_can_create_users": _can_create_ivg_users(current_igh_user),
        "igh_can_manage_users": bool(current_igh_user and (current_igh_user.is_superadmin() or current_igh_user.is_admin())),
        "igh_can_access_users_module": _can_access_users_module(current_igh_user),
        "igh_can_access_reports_module": _can_access_reports_module(current_igh_user),
        "igh_can_delete_records": _can_delete_operational_records(current_igh_user),
        "igh_can_edit_sales": _can_edit_ivg_sale(current_igh_user),
        "company_name": "Inversiones Guevara Herrera",
        "ivg_sale_categories": IVG_SALE_CATEGORIES,
        "ivg_category_labels": IVG_SALE_CATEGORY_LABELS,
        "ivg_category_metrics": category_metrics,
        "ivg_counts": {
            "users": _safe_count(IVGUser),
            "clients": _safe_count(IVGClient),
            "sales": _safe_count(IVGSale),
            "payments": _safe_count(IVGPayment),
            "cash_summaries": _safe_count(IVGCashSummary),
            "agenda": _safe_count(IVGAgendaItem),
            "credit_sales": _safe_count(IVGSale),
            "pending_invoices": _safe_count(IVGSale) if not _table_exists(IVGSale) else IVGSale.query.filter(IVGSale.balance_due > 0).count(),
        },
        "ivg_metrics": {
            "receivable_total": _safe_sum(IVGSale, IVGSale.balance_due) if _table_exists(IVGSale) else 0,
            "credit_total": _safe_sum(IVGSale, IVGSale.gross_amount) if _table_exists(IVGSale) else 0,
            "cash_total": _safe_sum(IVGCashSummary, IVGCashSummary.total_amount) if _table_exists(IVGCashSummary) else 0,
            "payment_total": _safe_sum(IVGPayment, IVGPayment.amount) if _table_exists(IVGPayment) else 0,
            "sales_month_total": (
                (_safe_sum(IVGSale, IVGSale.gross_amount, IVGSale.sale_date >= month_start) if _table_exists(IVGSale) else 0) +
                (_safe_sum(IVGCashSummary, IVGCashSummary.total_amount, IVGCashSummary.summary_date >= month_start) if _table_exists(IVGCashSummary) else 0)
            ),
            "collections_month_total": _safe_sum(IVGPayment, IVGPayment.amount, IVGPayment.payment_date >= month_start) if _table_exists(IVGPayment) else 0,
            "transfer_total": _safe_sum(IVGCashSummary, IVGCashSummary.transfer_amount) if _table_exists(IVGCashSummary) else 0,
            "efectivo_total": _safe_sum(IVGCashSummary, IVGCashSummary.cash_amount) if _table_exists(IVGCashSummary) else 0,
            "herbicidas_credito": _safe_sum(IVGSale, IVGSale.gross_amount, IVGSale.category == "herbicidas") if _table_exists(IVGSale) else 0,
            "concentrados_credito": _safe_sum(IVGSale, IVGSale.gross_amount, IVGSale.category == "concentrados") if _table_exists(IVGSale) else 0,
            "semillas_credito": _safe_sum(IVGSale, IVGSale.gross_amount, IVGSale.category == "semillas") if _table_exists(IVGSale) else 0,
            "latest_opening_amount": float(_to_decimal(latest_cash_summary.opening_amount)) if latest_cash_summary else 0,
            "latest_withdrawal_amount": float(_to_decimal(latest_cash_summary.withdrawal_amount)) if latest_cash_summary else 0,
            "latest_opening_plus_sales_amount": float(latest_opening_plus_sales) if latest_cash_summary else 0,
            "latest_expected_close_amount": float(_to_decimal(latest_cash_summary.expected_close_amount)) if latest_cash_summary else 0,
            "latest_actual_close_amount": float(_to_decimal(latest_cash_summary.actual_close_amount)) if latest_cash_summary else 0,
            "latest_variance_amount": float(abs(latest_variance)) if latest_cash_summary else 0,
            "latest_variance_signed": float(latest_variance) if latest_cash_summary else 0,
        },
        "pending_invoices": pending_invoices,
        "recent_collections": recent_collections,
        "recent_cash_summaries": recent_cash_summaries,
        "latest_cash_summary": latest_cash_summary,
        "top_clients": top_clients,
        "cashier_missing_dates": cashier_missing_dates,
        "cashier_missing_dates_count": len(cashier_missing_dates),
        "cashier_pending_agenda_count": pending_agenda_count,
        "ivg_tables": [
            "ivg_usuarios",
            "ivg_clientes",
            "ivg_ventas",
            "ivg_pagos",
            "ivg_contado",
            "ivg_agenda",
        ],
    }


def _render_section(title: str, eyebrow: str, description: str, cta: str):
    context = _base_context()
    context.update({
        "section_title": title,
        "section_eyebrow": eyebrow,
        "section_description": description,
        "section_cta": cta,
    })
    return render_template("igh/section.html", **context)


def _overdue_invoices_query(sort_key: str | None = None):
    if not _table_exists(IVGSale):
        return None

    query = IVGSale.query.outerjoin(IVGClient, IVGSale.client_id == IVGClient.id).filter(
        IVGSale.balance_due > 0,
        IVGSale.due_date.isnot(None),
        IVGSale.due_date < _today_start_utc(),
    )
    if sort_key == "amount_desc":
        return query.order_by(IVGSale.balance_due.desc(), IVGSale.due_date.asc(), IVGSale.id.desc())
    return query.order_by(IVGSale.due_date.asc(), IVGSale.sale_date.asc(), IVGSale.id.asc())


def _overdue_invoice_rows(sort_key: str | None = None):
    query = _overdue_invoices_query(sort_key)
    return query.all() if query is not None else []


def _build_sale_form_data(sale: IVGSale | None = None, form=None):
    if form is not None:
        return {
            "client_id": (form.get("client_id") or "").strip(),
            "category": _normalize_ivg_category(form.get("category")),
            "gross_amount": (form.get("gross_amount") or "").strip(),
            "reference_number": (form.get("reference_number") or "").strip(),
            "sale_date": (form.get("sale_date") or "").strip(),
            "due_date": (form.get("due_date") or "").strip(),
            "notes": (form.get("notes") or "").strip(),
        }

    if sale is not None:
        return {
            "client_id": str(sale.client_id or ""),
            "category": _normalize_ivg_category(sale.category),
            "gross_amount": str(sale.gross_amount or ""),
            "reference_number": sale.reference_number or "",
            "sale_date": _format_date_local(sale.sale_date),
            "due_date": _format_date_local(sale.due_date),
            "notes": sale.notes or "",
        }

    now = datetime.now().replace(second=0, microsecond=0)
    return {
        "client_id": "",
        "category": "herbicidas",
        "gross_amount": "",
        "reference_number": "",
        "sale_date": _format_date_local(now),
        "due_date": _format_date_local(now + timedelta(days=30)),
        "notes": "",
    }


def _build_user_form_data(user: IVGUser | None = None, form=None, actor: IVGUser | None = None):
    if form is not None:
        return {
            "username": (form.get("username") or "").strip(),
            "full_name": (form.get("full_name") or "").strip(),
            "email": (form.get("email") or "").strip(),
            "role": (
                "cajero"
                if actor and actor.is_admin() and not actor.is_superadmin()
                else (form.get("role") or "cajero")
            ),
            "is_active": bool(form.get("is_active")) if user else True,
        }

    return {
        "username": user.username if user else "",
        "full_name": user.full_name if user and user.full_name else "",
        "email": user.email if user and user.email else "",
        "role": user.role if user else "cajero",
        "is_active": user.is_active if user else True,
    }


def _get_ivg_user_or_404(user_id: int) -> IVGUser:
    if not _table_exists(IVGUser):
        abort(404)
    user = db.session.get(IVGUser, user_id)
    if user is None:
        abort(404)
    return user


def _get_ivg_client_or_404(client_id: int) -> IVGClient:
    if not _table_exists(IVGClient):
        abort(404)
    client = db.session.get(IVGClient, client_id)
    if client is None:
        abort(404)
    return client


def _get_ivg_sale_or_404(sale_id: int) -> IVGSale:
    if not _table_exists(IVGSale):
        abort(404)
    sale = db.session.get(IVGSale, sale_id)
    if sale is None:
        abort(404)
    return sale


def _get_ivg_cash_or_404(summary_id: int) -> IVGCashSummary:
    if not _table_exists(IVGCashSummary):
        abort(404)
    summary = db.session.get(IVGCashSummary, summary_id)
    if summary is None:
        abort(404)
    return summary


def _get_ivg_payment_or_404(payment_id: int) -> IVGPayment:
    if not _table_exists(IVGPayment):
        abort(404)
    payment = db.session.get(IVGPayment, payment_id)
    if payment is None:
        abort(404)
    return payment


def _get_ivg_agenda_or_404(item_id: int) -> IVGAgendaItem:
    if not _table_exists(IVGAgendaItem):
        abort(404)
    item = db.session.get(IVGAgendaItem, item_id)
    if item is None:
        abort(404)
    return item


@igh_bp.route("/")
@igh_login_required
def dashboard():
    context = _base_context()
    context["page_heading"] = "Dashboard"
    return render_template("igh/dashboard.html", **context)


@igh_bp.route("/dashboard")
@igh_login_required
def dashboard_alias():
    return redirect(url_for("igh.dashboard"))


@igh_bp.route("/usuarios")
@igh_login_required
@igh_admin_required
def usuarios():
    current_user = _current_igh_user()
    users = []
    if _table_exists(IVGUser):
        query = IVGUser.query.order_by(IVGUser.created_at.asc())
        users = [user for user in query.all() if _can_view_ivg_user(current_user, user)]
    context = _base_context()
    context["users"] = users
    return render_template("igh/users_list.html", **context)


@igh_bp.route("/usuarios/new", methods=["GET", "POST"])
@igh_login_required
@igh_admin_required
def usuarios_new():
    context = _base_context()
    actor = context["igh_current_user"]
    context["form_data"] = _build_user_form_data(actor=actor)

    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        full_name = (request.form.get("full_name") or "").strip()
        email = (request.form.get("email") or "").strip().lower() or None
        password = request.form.get("password") or ""
        role = "cajero" if actor and actor.is_admin() and not actor.is_superadmin() else (request.form.get("role") or "cajero")
        context["form_data"] = _build_user_form_data(form=request.form, actor=actor)

        if not username or not password:
            flash("Usuario y contraseña son obligatorios.", "danger")
            context["user"] = None
            return render_template("igh/users_form.html", **context)

        if len(password) < 4:
            flash("La contraseña debe tener al menos 4 caracteres.", "danger")
            context["user"] = None
            return render_template("igh/users_form.html", **context)

        if _table_exists(IVGUser):
            exists = IVGUser.query.filter_by(username=username).first()
            if exists:
                flash("Ya existe un usuario IVG con ese nombre.", "danger")
                context["user"] = None
                return render_template("igh/users_form.html", **context)

        user = IVGUser(
            username=username,
            full_name=full_name or None,
            email=email,
            role=role,
            is_active=True,
        )
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        flash(f"Usuario IVG {username} creado.", "success")
        return redirect(url_for("igh.usuarios"))

    context["user"] = None
    return render_template("igh/users_form.html", **context)


@igh_bp.route("/usuarios/<int:user_id>/edit", methods=["GET", "POST"])
@igh_login_required
def usuarios_edit(user_id: int):
    user = _get_ivg_user_or_404(user_id)
    context = _base_context()
    actor = context["igh_current_user"]

    if not _can_edit_ivg_user(actor, user):
        abort(403)

    if request.method == "POST":
        if actor and actor.is_superadmin():
            user.full_name = (request.form.get("full_name") or "").strip() or None
            user.email = (request.form.get("email") or "").strip().lower() or None
            user.role = request.form.get("role") or user.role
            user.is_active = bool(request.form.get("is_active"))

        new_password = request.form.get("password") or ""
        if new_password:
            if len(new_password) < 4:
                flash("La contraseña debe tener al menos 4 caracteres.", "danger")
                context["user"] = user
                context["form_data"] = _build_user_form_data(user=user, form=request.form, actor=actor)
                return render_template("igh/users_form.html", **context)
            user.set_password(new_password)

        db.session.commit()
        flash(
            (
                f"Clave de {user.username} actualizada."
                if actor and actor.is_admin() and not actor.is_superadmin()
                else f"Usuario IVG {user.username} actualizado."
            ),
            "success",
        )
        return redirect(url_for("igh.usuarios"))

    context["user"] = user
    context["form_data"] = _build_user_form_data(user=user, actor=actor)
    return render_template("igh/users_form.html", **context)


@igh_bp.route("/usuarios/<int:user_id>/delete", methods=["POST"])
@igh_login_required
@igh_admin_required
def usuarios_delete(user_id: int):
    user = _get_ivg_user_or_404(user_id)
    actor = _current_igh_user()
    if not _can_delete_ivg_user(actor, user):
        abort(403)

    username = user.username
    db.session.delete(user)
    db.session.commit()
    flash(f"Usuario IVG {username} eliminado.", "success")
    return redirect(url_for("igh.usuarios"))


@igh_bp.route("/clientes")
@igh_login_required
def clientes():
    search_query = (request.args.get("q") or "").strip()
    clients = []
    if _table_exists(IVGClient):
        query = IVGClient.query
        if search_query:
            term = f"%{search_query}%"
            query = query.filter(or_(
                IVGClient.name.ilike(term),
                IVGClient.legal_name.ilike(term),
                IVGClient.tax_id.ilike(term),
                IVGClient.email.ilike(term),
                IVGClient.phone.ilike(term),
                IVGClient.city.ilike(term),
            ))
        clients = query.order_by(IVGClient.created_at.desc()).all()
    for client in clients:
        client.can_delete = _client_delete_block_reason(client) is None
    context = _base_context()
    context["clients"] = clients
    context["search_query"] = search_query
    return render_template("igh/clients_list.html", **context)


@igh_bp.route("/clientes/<int:client_id>")
@igh_login_required
def clientes_detail(client_id: int):
    client = _get_ivg_client_or_404(client_id)
    sales = (
        IVGSale.query.filter_by(client_id=client.id)
        .order_by(IVGSale.sale_date.desc(), IVGSale.id.desc())
        .all()
        if _table_exists(IVGSale)
        else []
    )
    payments = (
        IVGPayment.query.join(IVGSale, IVGPayment.sale_id == IVGSale.id)
        .filter(IVGSale.client_id == client.id)
        .order_by(IVGPayment.payment_date.desc(), IVGPayment.id.desc())
        .all()
        if _table_exists(IVGPayment) and _table_exists(IVGSale)
        else []
    )
    pending_sales = [sale for sale in sales if Decimal(sale.balance_due or 0) > 0]
    paid_sales = [sale for sale in sales if Decimal(sale.balance_due or 0) <= 0]
    total_sales = sum(Decimal(sale.gross_amount or 0) for sale in sales)
    total_collections = sum(Decimal(payment.amount or 0) for payment in payments)
    pending_balance = sum(Decimal(sale.balance_due or 0) for sale in pending_sales)

    context = _base_context()
    context.update({
        "client": client,
        "client_sales": sales,
        "client_payments": payments,
        "client_can_delete": _client_delete_block_reason(client) is None,
        "client_delete_block_reason": _client_delete_block_reason(client),
        "client_pending_sales": pending_sales,
        "client_paid_sales": paid_sales,
        "client_metrics": {
            "sales_total": total_sales,
            "collections_total": total_collections,
            "pending_balance": pending_balance,
        },
    })
    return render_template("igh/client_detail.html", **context)


@igh_bp.route("/clientes/new", methods=["GET", "POST"])
@igh_login_required
def clientes_new():
    context = _base_context()
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        legal_name = (request.form.get("legal_name") or "").strip() or None
        tax_id = (request.form.get("tax_id") or "").strip() or None
        email = (request.form.get("email") or "").strip().lower() or None
        phone = (request.form.get("phone") or "").strip() or None
        city = (request.form.get("city") or "").strip() or None

        if not name:
            flash("El nombre del cliente es obligatorio.", "danger")
            return redirect(url_for("igh.clientes_new"))

        client = IVGClient(
            name=name,
            legal_name=legal_name,
            tax_id=tax_id,
            email=email,
            phone=phone,
            city=city,
            status="activo",
            is_active=True,
        )
        db.session.add(client)
        db.session.commit()
        flash(f"Cliente IVG {name} creado.", "success")
        return redirect(url_for("igh.clientes"))

    context["client"] = None
    return render_template("igh/clients_form.html", **context)


@igh_bp.route("/clientes/<int:client_id>/edit", methods=["GET", "POST"])
@igh_login_required
def clientes_edit(client_id: int):
    client = _get_ivg_client_or_404(client_id)
    context = _base_context()

    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("El nombre del cliente es obligatorio.", "danger")
            return redirect(url_for("igh.clientes_edit", client_id=client.id))

        client.name = name
        client.legal_name = (request.form.get("legal_name") or "").strip() or None
        client.tax_id = (request.form.get("tax_id") or "").strip() or None
        client.email = (request.form.get("email") or "").strip().lower() or None
        client.phone = (request.form.get("phone") or "").strip() or None
        client.city = (request.form.get("city") or "").strip() or None
        client.status = "activo"
        client.is_active = True

        db.session.commit()
        flash(f"Cliente IVG {client.name} actualizado.", "success")
        return redirect(url_for("igh.clientes"))

    context["client"] = client
    return render_template("igh/clients_form.html", **context)


@igh_bp.route("/clientes/<int:client_id>/delete", methods=["GET", "POST"])
@igh_login_required
@igh_admin_required
def clientes_delete(client_id: int):
    client = _get_ivg_client_or_404(client_id)
    block_reason = _client_delete_block_reason(client)

    if request.method == "GET":
        if block_reason:
            flash(block_reason, "warning")
            return redirect(url_for("igh.clientes_detail", client_id=client.id))

        context = _base_context()
        context.update({
            "client": client,
            "client_can_delete": True,
            "client_delete_block_reason": None,
        })
        return render_template("igh/client_delete_confirm.html", **context)

    if block_reason:
        flash(block_reason, "danger")
        return redirect(url_for("igh.clientes_detail", client_id=client.id))

    client_name = client.name
    try:
        db.session.delete(client)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash("No se pudo eliminar el cliente porque tiene movimientos relacionados.", "danger")
        return redirect(url_for("igh.clientes_detail", client_id=client.id))
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Error eliminando cliente IVG %s", client.id)
        flash("No se pudo eliminar el cliente por un error interno.", "danger")
        return redirect(url_for("igh.clientes_detail", client_id=client.id))

    flash(f"Cliente IVG {client_name} eliminado.", "success")
    return redirect(url_for("igh.clientes"))


@igh_bp.route("/ventas")
@igh_login_required
def ventas():
    search_query = (request.args.get("q") or "").strip()
    sales = []
    if _table_exists(IVGSale):
        query = IVGSale.query.outerjoin(IVGClient, IVGSale.client_id == IVGClient.id)
        if search_query:
            term = f"%{search_query}%"
            query = query.filter(or_(
                IVGSale.reference_number.ilike(term),
                IVGSale.category.ilike(term),
                IVGSale.status.ilike(term),
                IVGSale.notes.ilike(term),
                IVGClient.name.ilike(term),
                IVGClient.legal_name.ilike(term),
                IVGClient.tax_id.ilike(term),
            ))
        sales = query.order_by(IVGSale.sale_date.desc(), IVGSale.id.desc()).all()
    context = _base_context()
    context["sales"] = sales
    context["search_query"] = search_query
    return render_template("igh/sales_list.html", **context)


@igh_bp.route("/ventas/new", methods=["GET", "POST"])
@igh_login_required
def ventas_new():
    context = _base_context()
    context["clients"] = IVGClient.query.filter_by(is_active=True).order_by(IVGClient.name.asc()).all() if _table_exists(IVGClient) else []
    context["form_data"] = _build_sale_form_data()

    if request.method == "POST":
        context["form_data"] = _build_sale_form_data(form=request.form)
        try:
            client_id = int(request.form.get("client_id") or "0")
        except ValueError:
            client_id = 0
        category = _normalize_ivg_category(request.form.get("category"))
        reference_number = (request.form.get("reference_number") or "").strip() or None
        notes = (request.form.get("notes") or "").strip() or None

        try:
            gross_amount = _parse_decimal(request.form.get("gross_amount"), "El monto bruto")
            sale_date = _parse_datetime(request.form.get("sale_date"), "La fecha de venta", required=True)
            due_date = _parse_datetime(request.form.get("due_date"), "La fecha de vencimiento", required=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            context["sale"] = None
            return render_template("igh/sales_form.html", **context)

        client = db.session.get(IVGClient, client_id) if client_id else None
        if client is None:
            flash("Debes seleccionar un cliente válido.", "danger")
            context["sale"] = None
            return render_template("igh/sales_form.html", **context)
        duplicate_sale = _duplicate_sale_reference(reference_number)
        if duplicate_sale is not None:
            flash(f"La referencia {reference_number} ya existe en otra factura.", "danger")
            context["sale"] = None
            return render_template("igh/sales_form.html", **context)

        sale = IVGSale(
            client_id=client.id,
            sale_type="credito",
            category=category,
            gross_amount=gross_amount,
            balance_due=gross_amount,
            sale_date=sale_date,
            due_date=due_date,
            reference_number=reference_number,
            notes=notes,
            status="registrada",
        )
        _recalculate_sale_balance(sale)
        db.session.add(sale)
        db.session.commit()
        flash("Factura crédito IVG registrada.", "success")
        return redirect(url_for("igh.ventas"))

    context["sale"] = None
    return render_template("igh/sales_form.html", **context)


@igh_bp.route("/ventas/<int:sale_id>/edit", methods=["GET", "POST"])
@igh_login_required
def ventas_edit(sale_id: int):
    sale = _get_ivg_sale_or_404(sale_id)
    context = _base_context()
    if not _can_edit_ivg_sale(context["igh_current_user"]):
        abort(403)

    context["clients"] = IVGClient.query.filter_by(is_active=True).order_by(IVGClient.name.asc()).all() if _table_exists(IVGClient) else []
    context["form_data"] = _build_sale_form_data(sale=sale)

    if request.method == "POST":
        context["form_data"] = _build_sale_form_data(sale=sale, form=request.form)
        try:
            client_id = int(request.form.get("client_id") or "0")
        except ValueError:
            client_id = 0
        category = _normalize_ivg_category(request.form.get("category") or sale.category)
        reference_number = (request.form.get("reference_number") or "").strip() or None
        notes = (request.form.get("notes") or "").strip() or None

        try:
            gross_amount = _parse_decimal(request.form.get("gross_amount"), "El monto bruto")
            sale_date = _parse_datetime(request.form.get("sale_date"), "La fecha de venta", required=True)
            due_date = _parse_datetime(request.form.get("due_date"), "La fecha de vencimiento", required=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            context["sale"] = sale
            return render_template("igh/sales_form.html", **context)

        client = db.session.get(IVGClient, client_id) if client_id else None
        if client is None:
            flash("Debes seleccionar un cliente válido.", "danger")
            context["sale"] = sale
            return render_template("igh/sales_form.html", **context)
        duplicate_sale = _duplicate_sale_reference(reference_number, exclude_sale_id=sale.id)
        if duplicate_sale is not None:
            flash(f"La referencia {reference_number} ya existe en otra factura.", "danger")
            context["sale"] = sale
            return render_template("igh/sales_form.html", **context)

        sale.client_id = client.id
        sale.sale_type = "credito"
        sale.category = category
        sale.gross_amount = gross_amount
        sale.sale_date = sale_date
        sale.due_date = due_date
        sale.reference_number = reference_number
        sale.notes = notes
        _recalculate_sale_balance(sale)
        db.session.commit()
        flash("Factura crédito IVG actualizada.", "success")
        return redirect(url_for("igh.ventas"))

    context["sale"] = sale
    return render_template("igh/sales_form.html", **context)


@igh_bp.route("/ventas/<int:sale_id>/delete", methods=["POST"])
@igh_login_required
@igh_admin_required
def ventas_delete(sale_id: int):
    sale = _get_ivg_sale_or_404(sale_id)
    reference = sale.reference_number or f"FACT-{sale.id}"
    db.session.delete(sale)
    db.session.commit()
    flash(f"Factura crédito {reference} eliminada.", "success")
    return redirect(url_for("igh.ventas"))


@igh_bp.route("/facturas-por-cobrar")
@igh_login_required
def facturas_por_cobrar():
    sort_key = request.args.get("sort") or "oldest"
    if sort_key not in {"oldest", "amount_desc"}:
        sort_key = "oldest"
    overdue_sales = _overdue_invoice_rows(sort_key)
    today = datetime.utcnow().date()
    overdue_total = sum(_to_decimal(sale.balance_due) for sale in overdue_sales)
    context = _base_context()
    context.update({
        "overdue_sales": overdue_sales,
        "overdue_total": overdue_total,
        "overdue_today": today,
        "sort_key": sort_key,
    })
    return render_template("igh/receivables_list.html", **context)


@igh_bp.route("/facturas-por-cobrar/descargar/<fmt>")
@igh_login_required
def facturas_por_cobrar_descargar(fmt: str):
    sort_key = request.args.get("sort") or "oldest"
    if sort_key not in {"oldest", "amount_desc"}:
        sort_key = "oldest"
    today_date = datetime.utcnow().date()
    sales = _overdue_invoice_rows(sort_key)
    rows = [
        [
            sale.reference_number or f"FACT-{sale.id}",
            sale.client.name if sale.client else "",
            _ivg_category_label(sale.category),
            sale.status,
            _format_report_date(sale.sale_date),
            _format_report_date(sale.due_date),
            max((today_date - sale.due_date.date()).days, 0) if sale.due_date else 0,
            _money_value(sale.gross_amount),
            _money_value(sale.balance_due),
            sale.notes or "",
        ]
        for sale in sales
    ]
    title = "Facturas por cobrar vencidas"
    headers = ["Factura", "Cliente", "Categoria", "Estado", "Fecha venta", "Vencimiento", "Dias vencida", "Total", "Saldo", "Notas"]
    file_date = datetime.utcnow().strftime("%Y%m%d")
    if fmt == "xlsx":
        return _xlsx_response(f"ivg-facturas-por-cobrar-{file_date}.xlsx", title, headers, rows)
    if fmt == "pdf":
        return _pdf_response(f"ivg-facturas-por-cobrar-{file_date}.pdf", title, headers, rows)
    abort(404)


@igh_bp.route("/contado")
@igh_login_required
def contado():
    summaries = IVGCashSummary.query.order_by(IVGCashSummary.summary_date.desc(), IVGCashSummary.id.desc()).all() if _table_exists(IVGCashSummary) else []
    context = _base_context()
    context["summaries"] = summaries
    return render_template("igh/cash_list.html", **context)


@igh_bp.route("/contado/new", methods=["GET", "POST"])
@igh_login_required
def contado_new():
    context = _base_context()
    previous_summary = _latest_cash_summary()
    context["default_opening_amount"] = previous_summary.actual_close_amount if previous_summary else Decimal("0.00")
    context["previous_summary"] = previous_summary

    if request.method == "POST":
        notes = (request.form.get("notes") or "").strip() or None
        try:
            summary_date = _parse_datetime(request.form.get("summary_date"), "La fecha del cierre", required=True)
            opening_amount = _parse_decimal(request.form.get("opening_amount"), "La apertura")
            cash_amount = _parse_decimal(request.form.get("cash_amount"), "La venta del día")
            withdrawal_amount = _parse_decimal(request.form.get("withdrawal_amount"), "El retiro en efectivo")
            transfer_amount = _parse_decimal(request.form.get("transfer_amount"), "Las transferencias")
            actual_close_amount = _parse_decimal(request.form.get("actual_close_amount"), "El cierre real")
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("igh.contado_new"))

        close_values = _calculate_day_close(opening_amount, cash_amount, withdrawal_amount, transfer_amount)
        total_amount = close_values["sales_total"]
        if total_amount <= 0:
            flash("Debes registrar al menos una venta o transferencia mayor que cero.", "danger")
            return redirect(url_for("igh.contado_new"))
        expected_close_amount = close_values["expected_close_amount"]
        variance_amount = actual_close_amount - expected_close_amount

        summary = IVGCashSummary(
            summary_date=summary_date,
            opening_amount=opening_amount,
            cash_amount=cash_amount,
            withdrawal_amount=withdrawal_amount,
            transfer_amount=transfer_amount,
            total_amount=total_amount,
            expected_close_amount=expected_close_amount,
            actual_close_amount=actual_close_amount,
            variance_amount=variance_amount,
            notes=notes,
        )
        db.session.add(summary)
        db.session.commit()
        flash("Cierre del día registrado.", "success")
        return redirect(url_for("igh.contado"))

    context["summary"] = None
    return render_template("igh/cash_form.html", **context)


@igh_bp.route("/contado/<int:summary_id>/edit", methods=["GET", "POST"])
@igh_login_required
def contado_edit(summary_id: int):
    summary = _get_ivg_cash_or_404(summary_id)
    context = _base_context()

    if request.method == "POST":
        notes = (request.form.get("notes") or "").strip() or None
        try:
            summary_date = _parse_datetime(request.form.get("summary_date"), "La fecha del cierre", required=True)
            opening_amount = _parse_decimal(request.form.get("opening_amount"), "La apertura")
            cash_amount = _parse_decimal(request.form.get("cash_amount"), "La venta del día")
            withdrawal_amount = _parse_decimal(request.form.get("withdrawal_amount"), "El retiro en efectivo")
            transfer_amount = _parse_decimal(request.form.get("transfer_amount"), "Las transferencias")
            actual_close_amount = _parse_decimal(request.form.get("actual_close_amount"), "El cierre real")
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("igh.contado_edit", summary_id=summary.id))

        close_values = _calculate_day_close(opening_amount, cash_amount, withdrawal_amount, transfer_amount)
        total_amount = close_values["sales_total"]
        if total_amount <= 0:
            flash("Debes registrar al menos una venta o transferencia mayor que cero.", "danger")
            return redirect(url_for("igh.contado_edit", summary_id=summary.id))
        expected_close_amount = close_values["expected_close_amount"]
        variance_amount = actual_close_amount - expected_close_amount

        summary.summary_date = summary_date
        summary.opening_amount = opening_amount
        summary.cash_amount = cash_amount
        summary.withdrawal_amount = withdrawal_amount
        summary.transfer_amount = transfer_amount
        summary.total_amount = total_amount
        summary.expected_close_amount = expected_close_amount
        summary.actual_close_amount = actual_close_amount
        summary.variance_amount = variance_amount
        summary.notes = notes
        db.session.commit()
        flash("Cierre del día actualizado.", "success")
        return redirect(url_for("igh.contado"))

    context["summary"] = summary
    return render_template("igh/cash_form.html", **context)


@igh_bp.route("/contado/<int:summary_id>/delete", methods=["POST"])
@igh_login_required
@igh_admin_required
def contado_delete(summary_id: int):
    summary = _get_ivg_cash_or_404(summary_id)
    db.session.delete(summary)
    db.session.commit()
    flash("Cierre del día eliminado.", "success")
    return redirect(url_for("igh.contado"))


@igh_bp.route("/pagos")
@igh_login_required
def pagos():
    search_query = (request.args.get("q") or "").strip()
    payments = []
    if _table_exists(IVGPayment):
        query = (
            IVGPayment.query
            .outerjoin(IVGSale, IVGPayment.sale_id == IVGSale.id)
            .outerjoin(IVGClient, IVGSale.client_id == IVGClient.id)
        )
        if search_query:
            term = f"%{search_query}%"
            query = query.filter(or_(
                IVGPayment.payment_kind.ilike(term),
                IVGPayment.payment_method.ilike(term),
                IVGPayment.category.ilike(term),
                IVGPayment.notes.ilike(term),
                IVGSale.reference_number.ilike(term),
                IVGSale.category.ilike(term),
                IVGClient.name.ilike(term),
                IVGClient.legal_name.ilike(term),
                IVGClient.tax_id.ilike(term),
            ))
        payments = query.order_by(IVGPayment.payment_date.desc(), IVGPayment.id.desc()).all()
    context = _base_context()
    context["payments"] = payments
    context["search_query"] = search_query
    return render_template("igh/payments_list.html", **context)


@igh_bp.route("/pagos/new", methods=["GET", "POST"])
@igh_login_required
def pagos_new():
    context = _base_context()
    context["credit_sales"] = IVGSale.query.filter(IVGSale.balance_due > 0).order_by(IVGSale.sale_date.desc()).all() if _table_exists(IVGSale) else []
    selected_sale_id = (request.args.get("sale_id") or "").strip()
    context["form_data"] = {
        "sale_id": selected_sale_id,
        "payment_kind": "abono",
        "payment_method": "",
        "amount": "",
        "payment_date": _format_date_local(datetime.utcnow()),
        "notes": "",
    }

    if request.method == "POST":
        context["form_data"] = {
            "sale_id": request.form.get("sale_id") or "",
            "payment_kind": (request.form.get("payment_kind") or "abono").strip(),
            "payment_method": (request.form.get("payment_method") or "").strip(),
            "amount": request.form.get("amount") or "",
            "payment_date": request.form.get("payment_date") or "",
            "notes": request.form.get("notes") or "",
        }
        try:
            sale_id = int(request.form.get("sale_id") or "0")
        except ValueError:
            sale_id = 0
        payment_kind = (request.form.get("payment_kind") or "abono").strip()
        payment_method = (request.form.get("payment_method") or "").strip()
        notes = (request.form.get("notes") or "").strip() or None

        try:
            amount = _parse_decimal(request.form.get("amount"), "El monto")
            payment_date = _parse_datetime(request.form.get("payment_date"), "La fecha de pago", required=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            context["payment"] = None
            return render_template("igh/payments_form.html", **context)

        sale = db.session.get(IVGSale, sale_id) if sale_id else None
        if sale is None:
            flash("Debes seleccionar una factura de crédito válida.", "danger")
            context["payment"] = None
            return render_template("igh/payments_form.html", **context)
        if amount <= 0:
            flash("El monto debe ser mayor que cero.", "danger")
            context["payment"] = None
            return render_template("igh/payments_form.html", **context)
        if Decimal(sale.balance_due or 0) <= 0:
            flash("La factura seleccionada ya no tiene saldo pendiente.", "warning")
            context["payment"] = None
            return render_template("igh/payments_form.html", **context)
        if amount > Decimal(sale.balance_due or 0):
            flash("El monto supera el saldo pendiente de la factura.", "danger")
            context["payment"] = None
            return render_template("igh/payments_form.html", **context)

        payment = IVGPayment(
            sale_id=sale.id,
            payment_kind=payment_kind,
            payment_method=payment_method,
            category=_normalize_ivg_category(sale.category),
            amount=amount,
            payment_date=payment_date,
            notes=notes,
        )
        db.session.add(payment)
        db.session.flush()
        _recalculate_sale_balance(sale)
        db.session.commit()
        flash("Pago / abono registrado.", "success")
        return redirect(url_for("igh.pagos"))

    context["payment"] = None
    return render_template("igh/payments_form.html", **context)


@igh_bp.route("/pagos/<int:payment_id>/edit", methods=["GET", "POST"])
@igh_login_required
def pagos_edit(payment_id: int):
    payment = _get_ivg_payment_or_404(payment_id)
    old_sale = payment.sale
    context = _base_context()
    context["credit_sales"] = IVGSale.query.order_by(IVGSale.sale_date.desc()).all() if _table_exists(IVGSale) else []
    context["form_data"] = {
        "sale_id": str(payment.sale_id or ""),
        "payment_kind": payment.payment_kind or "abono",
        "payment_method": payment.payment_method or "",
        "amount": payment.amount if payment.amount is not None else "",
        "payment_date": _format_date_local(payment.payment_date),
        "notes": payment.notes or "",
    }

    if request.method == "POST":
        context["form_data"] = {
            "sale_id": request.form.get("sale_id") or "",
            "payment_kind": (request.form.get("payment_kind") or payment.payment_kind).strip(),
            "payment_method": (request.form.get("payment_method") or payment.payment_method).strip(),
            "amount": request.form.get("amount") or "",
            "payment_date": request.form.get("payment_date") or "",
            "notes": request.form.get("notes") or "",
        }
        try:
            sale_id = int(request.form.get("sale_id") or "0")
        except ValueError:
            sale_id = 0
        payment_kind = (request.form.get("payment_kind") or payment.payment_kind).strip()
        payment_method = (request.form.get("payment_method") or payment.payment_method).strip()
        notes = (request.form.get("notes") or "").strip() or None

        try:
            amount = _parse_decimal(request.form.get("amount"), "El monto")
            payment_date = _parse_datetime(request.form.get("payment_date"), "La fecha de pago", required=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            context["payment"] = payment
            return render_template("igh/payments_form.html", **context)

        sale = db.session.get(IVGSale, sale_id) if sale_id else None
        if sale is None:
            flash("Debes seleccionar una factura de crédito válida.", "danger")
            context["payment"] = payment
            return render_template("igh/payments_form.html", **context)
        if amount <= 0:
            flash("El monto debe ser mayor que cero.", "danger")
            context["payment"] = payment
            return render_template("igh/payments_form.html", **context)

        payment.sale_id = sale.id
        payment.payment_kind = payment_kind
        payment.payment_method = payment_method
        payment.category = _normalize_ivg_category(sale.category)
        payment.amount = amount
        payment.payment_date = payment_date
        payment.notes = notes
        db.session.flush()
        _recalculate_sale_balance(old_sale)
        if sale.id != old_sale.id:
            _recalculate_sale_balance(sale)
        else:
            _recalculate_sale_balance(payment.sale)
        if Decimal(payment.sale.balance_due or 0) < 0:
            flash("El monto editado deja saldo negativo en la factura.", "danger")
            db.session.rollback()
            return redirect(url_for("igh.pagos_edit", payment_id=payment.id))
        db.session.commit()
        flash("Pago / abono actualizado.", "success")
        return redirect(url_for("igh.pagos"))

    context["payment"] = payment
    return render_template("igh/payments_form.html", **context)


@igh_bp.route("/pagos/<int:payment_id>/delete", methods=["POST"])
@igh_login_required
@igh_admin_required
def pagos_delete(payment_id: int):
    payment = _get_ivg_payment_or_404(payment_id)
    sale = payment.sale
    db.session.delete(payment)
    db.session.flush()
    if sale is not None:
        _recalculate_sale_balance(sale)
    db.session.commit()
    flash("Pago / abono eliminado.", "success")
    return redirect(url_for("igh.pagos"))


@igh_bp.route("/agenda")
@igh_login_required
def agenda():
    items = IVGAgendaItem.query.order_by(IVGAgendaItem.scheduled_for.asc(), IVGAgendaItem.id.desc()).all() if _table_exists(IVGAgendaItem) else []
    context = _base_context()
    context["agenda_items"] = items
    return render_template("igh/agenda_list.html", **context)


@igh_bp.route("/agenda/new", methods=["GET", "POST"])
@igh_login_required
def agenda_new():
    context = _base_context()
    context["clients"] = IVGClient.query.filter_by(is_active=True).order_by(IVGClient.name.asc()).all() if _table_exists(IVGClient) else []

    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        activity_type = (request.form.get("activity_type") or "seguimiento").strip()
        status = (request.form.get("status") or "pendiente").strip()
        priority = (request.form.get("priority") or "media").strip()
        notes = (request.form.get("notes") or "").strip() or None
        try:
            client_id = int(request.form.get("client_id") or "0")
        except ValueError:
            client_id = 0

        try:
            scheduled_for = _parse_datetime(request.form.get("scheduled_for"), "La fecha programada", required=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("igh.agenda_new"))

        if not title:
            flash("El título de la actividad es obligatorio.", "danger")
            return redirect(url_for("igh.agenda_new"))

        item = IVGAgendaItem(
            client_id=client_id or None,
            title=title,
            activity_type=activity_type,
            status=status,
            priority=priority,
            scheduled_for=scheduled_for,
            notes=notes,
        )
        db.session.add(item)
        db.session.commit()
        flash("Actividad de agenda creada.", "success")
        return redirect(url_for("igh.agenda"))

    context["item"] = None
    return render_template("igh/agenda_form.html", **context)


@igh_bp.route("/agenda/<int:item_id>/edit", methods=["GET", "POST"])
@igh_login_required
def agenda_edit(item_id: int):
    item = _get_ivg_agenda_or_404(item_id)
    context = _base_context()
    context["clients"] = IVGClient.query.filter_by(is_active=True).order_by(IVGClient.name.asc()).all() if _table_exists(IVGClient) else []

    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        activity_type = (request.form.get("activity_type") or item.activity_type).strip()
        status = (request.form.get("status") or item.status).strip()
        priority = (request.form.get("priority") or item.priority).strip()
        notes = (request.form.get("notes") or "").strip() or None
        try:
            client_id = int(request.form.get("client_id") or "0")
        except ValueError:
            client_id = 0

        try:
            scheduled_for = _parse_datetime(request.form.get("scheduled_for"), "La fecha programada", required=True)
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("igh.agenda_edit", item_id=item.id))

        if not title:
            flash("El título de la actividad es obligatorio.", "danger")
            return redirect(url_for("igh.agenda_edit", item_id=item.id))

        item.client_id = client_id or None
        item.title = title
        item.activity_type = activity_type
        item.status = status
        item.priority = priority
        item.scheduled_for = scheduled_for
        item.notes = notes
        db.session.commit()
        flash("Actividad de agenda actualizada.", "success")
        return redirect(url_for("igh.agenda"))

    context["item"] = item
    return render_template("igh/agenda_form.html", **context)


@igh_bp.route("/agenda/<int:item_id>/delete", methods=["POST"])
@igh_login_required
@igh_admin_required
def agenda_delete(item_id: int):
    item = _get_ivg_agenda_or_404(item_id)
    db.session.delete(item)
    db.session.commit()
    flash("Actividad eliminada.", "success")
    return redirect(url_for("igh.agenda"))


@igh_bp.route("/reportes")
@igh_login_required
@igh_admin_required
def reportes():
    context = _base_context()
    context["recent_sales"] = IVGSale.query.order_by(IVGSale.sale_date.desc(), IVGSale.id.desc()).limit(8).all() if _table_exists(IVGSale) else []
    context["recent_payments"] = IVGPayment.query.order_by(IVGPayment.payment_date.desc(), IVGPayment.id.desc()).limit(8).all() if _table_exists(IVGPayment) else []
    context["pending_agenda"] = IVGAgendaItem.query.filter(IVGAgendaItem.status != "completada").order_by(IVGAgendaItem.scheduled_for.asc()).limit(6).all() if _table_exists(IVGAgendaItem) else []
    sales_status_rows = []
    receivable_category_rows = []
    if _table_exists(IVGSale):
        sales_status_rows = (
            db.session.query(IVGSale.status, func.count(IVGSale.id))
            .group_by(IVGSale.status)
            .order_by(func.count(IVGSale.id).desc())
            .all()
        )
        receivable_category_rows = (
            db.session.query(IVGSale.category, func.coalesce(func.sum(IVGSale.balance_due), 0))
            .filter(IVGSale.balance_due > 0)
            .group_by(IVGSale.category)
            .order_by(func.coalesce(func.sum(IVGSale.balance_due), 0).desc())
            .all()
        )
    context["report_chart_data"] = {
        "payment_methods": [
            {"label": "Efectivo", "value": context["ivg_metrics"]["efectivo_total"]},
            {"label": "Transferencias", "value": context["ivg_metrics"]["transfer_total"]},
        ],
        "collection_categories": [
            {"label": item["label"], "value": item["collections_total"]}
            for item in context["ivg_category_metrics"]
        ],
        "sale_status": [
            {"label": status or "Sin estado", "value": count}
            for status, count in sales_status_rows
        ],
        "receivable_categories": [
            {"label": _ivg_category_label(category), "value": float(_to_decimal(total))}
            for category, total in receivable_category_rows
        ],
    }
    context["download_reports"] = [
        ("cierres", "Resumen de cierres", "Apertura, ventas, retiros, transferencias y diferencias."),
        ("ventas-contado", "Ventas de contado", "Venta diaria separada por efectivo y transferencia."),
        ("ventas-credito", "Ventas de crédito", "Facturas a crédito con cliente, categoría, total y saldo."),
        ("cobros", "Cobros registrados", "Pagos y abonos aplicados a facturas de crédito."),
        ("cartera", "Cartera por cobrar", "Facturas pendientes, parciales y vencimientos."),
        ("clientes-saldos", "Clientes con saldo", "Saldo pendiente acumulado por cliente."),
        ("agenda", "Agenda pendiente", "Seguimientos y actividades no completadas."),
    ]
    return render_template("igh/reports.html", **context)


def _ivg_report_payload(report_type: str):
    if report_type == "cierres":
        summaries = IVGCashSummary.query.order_by(IVGCashSummary.summary_date.desc(), IVGCashSummary.id.desc()).all() if _table_exists(IVGCashSummary) else []
        rows = [
            [
                _format_report_date(summary.summary_date),
                _money_value(summary.opening_amount),
                _money_value(summary.cash_amount),
                _money_value(_to_decimal(summary.opening_amount) + _to_decimal(summary.cash_amount)),
                _money_value(summary.withdrawal_amount),
                _money_value(summary.transfer_amount),
                _money_value(summary.expected_close_amount),
                _money_value(summary.actual_close_amount),
                _money_value(summary.variance_amount),
                summary.notes or "",
            ]
            for summary in summaries
        ]
        return (
            "ivg-resumen-cierres",
            "Resumen de cierres",
            ["Fecha", "Apertura", "Venta efectivo", "Venta efectivo + apertura", "Retiro efectivo", "Transferencias", "Cierre esperado", "Cierre real", "Diferencia", "Notas"],
            rows,
        )

    if report_type == "ventas-contado":
        summaries = IVGCashSummary.query.order_by(IVGCashSummary.summary_date.desc(), IVGCashSummary.id.desc()).all() if _table_exists(IVGCashSummary) else []
        rows = [
            [
                _format_report_date(summary.summary_date),
                _money_value(summary.cash_amount),
                _money_value(summary.transfer_amount),
                _money_value(summary.total_amount),
                _money_value(summary.withdrawal_amount),
                _money_value(summary.actual_close_amount),
                summary.notes or "",
            ]
            for summary in summaries
        ]
        return (
            "ivg-ventas-contado",
            "Ventas de contado",
            ["Fecha", "Efectivo", "Transferencias", "Venta total", "Retiro efectivo", "Cierre real", "Notas"],
            rows,
        )

    if report_type == "ventas-credito":
        sales = IVGSale.query.order_by(IVGSale.sale_date.desc(), IVGSale.id.desc()).all() if _table_exists(IVGSale) else []
        rows = [
            [
                _format_report_date(sale.sale_date),
                sale.reference_number or f"FACT-{sale.id}",
                sale.client.name if sale.client else "",
                _ivg_category_label(sale.category),
                sale.status,
                _format_report_date(sale.due_date),
                _money_value(sale.gross_amount),
                _money_value(sale.balance_due),
                sale.notes or "",
            ]
            for sale in sales
        ]
        return (
            "ivg-ventas-credito",
            "Ventas de credito",
            ["Fecha", "Referencia", "Cliente", "Categoria", "Estado", "Vencimiento", "Total", "Saldo", "Notas"],
            rows,
        )

    if report_type == "cobros":
        payments = IVGPayment.query.order_by(IVGPayment.payment_date.desc(), IVGPayment.id.desc()).all() if _table_exists(IVGPayment) else []
        rows = [
            [
                _format_report_date(payment.payment_date),
                payment.sale.reference_number or f"FACT-{payment.sale.id}" if payment.sale else "",
                payment.sale.client.name if payment.sale and payment.sale.client else "",
                payment.payment_kind,
                payment.payment_method,
                _ivg_category_label(payment.category),
                _money_value(payment.amount),
                payment.notes or "",
            ]
            for payment in payments
        ]
        return (
            "ivg-cobros",
            "Cobros registrados",
            ["Fecha", "Factura", "Cliente", "Tipo", "Metodo", "Categoria", "Monto", "Notas"],
            rows,
        )

    if report_type == "cartera":
        sales = IVGSale.query.filter(IVGSale.balance_due > 0).order_by(IVGSale.due_date.asc(), IVGSale.sale_date.desc()).all() if _table_exists(IVGSale) else []
        rows = [
            [
                sale.reference_number or f"FACT-{sale.id}",
                sale.client.name if sale.client else "",
                _ivg_category_label(sale.category),
                sale.status,
                _format_report_date(sale.sale_date),
                _format_report_date(sale.due_date),
                _money_value(sale.gross_amount),
                _money_value(sale.balance_due),
            ]
            for sale in sales
        ]
        return (
            "ivg-cartera-por-cobrar",
            "Cartera por cobrar",
            ["Factura", "Cliente", "Categoria", "Estado", "Fecha venta", "Vencimiento", "Total", "Saldo"],
            rows,
        )

    if report_type == "clientes-saldos":
        rows = []
        if _table_exists(IVGClient) and _table_exists(IVGSale):
            data = (
                db.session.query(
                    IVGClient.name,
                    func.count(IVGSale.id),
                    func.coalesce(func.sum(IVGSale.gross_amount), 0),
                    func.coalesce(func.sum(IVGSale.balance_due), 0),
                )
                .join(IVGSale, IVGSale.client_id == IVGClient.id)
                .filter(IVGSale.balance_due > 0)
                .group_by(IVGClient.id, IVGClient.name)
                .order_by(func.coalesce(func.sum(IVGSale.balance_due), 0).desc())
                .all()
            )
            rows = [[name, count, _money_value(total), _money_value(balance)] for name, count, total, balance in data]
        return (
            "ivg-clientes-con-saldo",
            "Clientes con saldo",
            ["Cliente", "Facturas pendientes", "Total facturado", "Saldo pendiente"],
            rows,
        )

    if report_type == "agenda":
        items = IVGAgendaItem.query.filter(IVGAgendaItem.status != "completada").order_by(IVGAgendaItem.scheduled_for.asc()).all() if _table_exists(IVGAgendaItem) else []
        rows = [
            [
                _format_report_date(item.scheduled_for),
                item.client.name if item.client else "",
                item.title,
                item.activity_type,
                item.priority,
                item.status,
                item.notes or "",
            ]
            for item in items
        ]
        return (
            "ivg-agenda-pendiente",
            "Agenda pendiente",
            ["Fecha", "Cliente", "Actividad", "Tipo", "Prioridad", "Estado", "Notas"],
            rows,
        )

    abort(404)


@igh_bp.route("/reportes/descargar/<report_type>/<fmt>")
@igh_login_required
@igh_admin_required
def reportes_descargar(report_type: str, fmt: str):
    today = datetime.utcnow().strftime("%Y%m%d")
    slug, title, headers, rows = _ivg_report_payload(report_type)
    if fmt == "xlsx":
        return _xlsx_response(f"{slug}-{today}.xlsx", title, headers, rows)
    if fmt == "pdf":
        return _pdf_response(f"{slug}-{today}.pdf", title, headers, rows)
    abort(404)


@igh_bp.route("/productos")
@igh_login_required
def productos_alias():
    return redirect(url_for("igh.ventas"))


@igh_bp.route("/cobros")
@igh_login_required
def cobros_alias():
    return redirect(url_for("igh.pagos"))


@igh_bp.route("/logout")
def logout():
    session.pop("igh_user_id", None)
    session.pop("igh_user", None)
    session.pop("igh_role", None)
    session.pop("tenant_id", None)
    flash("Sesión de Inversiones Guevara Herrera cerrada.", "info")
    return redirect(url_for("auth.login"))
