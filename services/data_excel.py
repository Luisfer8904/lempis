"""
Importación y exportación de datos base del tenant en Excel.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models import db
from models.catalog import Category, Customer, Product
from models.suppliers import Supplier


PRODUCT_HEADERS = [
    "sku", "nombre", "descripcion", "categoria", "tipo", "precio_general",
    "precio_mayorista", "precio_especial", "costo", "stock", "controla_stock",
    "controla_lotes", "activo", "imagen_url",
]
CUSTOMER_HEADERS = [
    "nombre", "rtn", "email", "telefono", "direccion", "ciudad", "tarifa",
    "pais", "notas", "activo",
]
SUPPLIER_HEADERS = [
    "codigo", "nombre", "razon_social", "rtn", "contacto", "email", "telefono",
    "direccion", "ciudad", "pais", "dias_credito", "limite_credito",
    "precios_incluyen_impuesto", "notas", "activo",
]
CATEGORY_HEADERS = ["nombre", "descripcion", "color"]


def build_export_workbook(tenant) -> BytesIO:
    wb = Workbook()
    _write_categories(wb.active, tenant)
    _write_products(wb.create_sheet("Productos"), tenant)
    _write_customers(wb.create_sheet("Clientes"), tenant)
    _write_suppliers(wb.create_sheet("Proveedores"), tenant)
    return _to_stream(wb)


def build_template_workbook() -> BytesIO:
    wb = Workbook()
    _setup_sheet(wb.active, "Categorias", CATEGORY_HEADERS)
    _setup_sheet(wb.create_sheet("Productos"), "Productos", PRODUCT_HEADERS)
    _setup_sheet(wb.create_sheet("Clientes"), "Clientes", CUSTOMER_HEADERS)
    _setup_sheet(wb.create_sheet("Proveedores"), "Proveedores", SUPPLIER_HEADERS)

    wb["Categorias"].append(["Concentrados", "Alimentos y suplementos", "indigo"])
    wb["Productos"].append([
        "PROD-001", "Producto ejemplo", "Descripcion opcional", "Concentrados",
        "product", 100, 95, 90, 70, 10, "si", "si", "si", "",
    ])
    wb["Clientes"].append([
        "Cliente ejemplo", "08011990000000", "cliente@example.com", "9999-9999",
        "Direccion", "Catacamas", "general", "HN", "", "si",
    ])
    wb["Proveedores"].append([
        "P-0001", "Proveedor ejemplo", "Proveedor Ejemplo S.A.", "08011990000001",
        "Contacto", "proveedor@example.com", "8888-8888", "Direccion", "Catacamas",
        "HN", 30, 0, "no", "", "si",
    ])
    for ws in wb.worksheets:
        _autosize(ws)
    return _to_stream(wb)


def import_workbook(tenant, file_storage) -> dict:
    wb = load_workbook(file_storage, data_only=True)
    stats = {
        "categories_created": 0,
        "categories_updated": 0,
        "products_created": 0,
        "products_updated": 0,
        "customers_created": 0,
        "customers_updated": 0,
        "suppliers_created": 0,
        "suppliers_updated": 0,
        "errors": [],
    }

    category_by_name = {
        c.name.strip().lower(): c
        for c in Category.query.filter_by(tenant_id=tenant.id).all()
    }

    if "Categorias" in wb.sheetnames:
        _import_categories(wb["Categorias"], tenant, category_by_name, stats)
    if "Productos" in wb.sheetnames:
        _import_products(wb["Productos"], tenant, category_by_name, stats)
    if "Clientes" in wb.sheetnames:
        _import_customers(wb["Clientes"], tenant, stats)
    if "Proveedores" in wb.sheetnames:
        _import_suppliers(wb["Proveedores"], tenant, stats)

    db.session.commit()
    return stats


def _write_categories(ws, tenant) -> None:
    _setup_sheet(ws, "Categorias", CATEGORY_HEADERS)
    for cat in Category.query.filter_by(tenant_id=tenant.id).order_by(Category.name.asc()).all():
        ws.append([cat.name, cat.description, cat.color])
    _autosize(ws)


def _write_products(ws, tenant) -> None:
    _setup_sheet(ws, "Productos", PRODUCT_HEADERS)
    rows = Product.query.filter_by(tenant_id=tenant.id).order_by(Product.name.asc()).all()
    for p in rows:
        ws.append([
            p.sku, p.name, p.description, p.category.name if p.category else "",
            p.kind, p.price, p.price_wholesale, p.price_special, p.cost, p.stock,
            _yes_no(p.track_stock), _yes_no(p.track_batches), _yes_no(p.is_active),
            p.image_url,
        ])
    _autosize(ws)


def _write_customers(ws, tenant) -> None:
    _setup_sheet(ws, "Clientes", CUSTOMER_HEADERS)
    rows = Customer.query.filter_by(tenant_id=tenant.id).order_by(Customer.name.asc()).all()
    for c in rows:
        ws.append([
            c.name, c.tax_id, c.email, c.phone, c.address, c.city,
            c.preferred_price_tier, c.country_code, c.notes, _yes_no(c.is_active),
        ])
    _autosize(ws)


def _write_suppliers(ws, tenant) -> None:
    _setup_sheet(ws, "Proveedores", SUPPLIER_HEADERS)
    rows = Supplier.query.filter_by(tenant_id=tenant.id).order_by(Supplier.name.asc()).all()
    for s in rows:
        ws.append([
            s.code, s.name, s.legal_name, s.tax_id, s.contact_name, s.email, s.phone,
            s.address, s.city, s.country_code, s.payment_terms_days, s.credit_limit,
            _yes_no(s.price_includes_tax), s.notes, _yes_no(s.is_active),
        ])
    _autosize(ws)


def _import_categories(ws, tenant, category_by_name, stats) -> None:
    for rownum, row in _rows(ws, CATEGORY_HEADERS, "Categorias", stats):
        name = _text(row.get("nombre"))
        if not name:
            _error(stats, f"Categorias fila {rownum}: nombre es obligatorio.")
            continue
        key = name.lower()
        cat = category_by_name.get(key)
        created = cat is None
        if created:
            cat = Category(tenant_id=tenant.id, name=name)
            db.session.add(cat)
            category_by_name[key] = cat
        cat.description = _text(row.get("descripcion")) or None
        cat.color = _text(row.get("color")) or None
        stats["categories_created" if created else "categories_updated"] += 1


def _import_products(ws, tenant, category_by_name, stats) -> None:
    for rownum, row in _rows(ws, PRODUCT_HEADERS, "Productos", stats):
        sku = _text(row.get("sku"))
        name = _text(row.get("nombre"))
        if not sku or not name:
            _error(stats, f"Productos fila {rownum}: sku y nombre son obligatorios.")
            continue

        product = Product.query.filter_by(tenant_id=tenant.id, sku=sku).first()
        created = product is None
        if created:
            product = Product(tenant_id=tenant.id, sku=sku)
            db.session.add(product)

        product.name = name
        product.description = _text(row.get("descripcion")) or None
        product.kind = _choice(row.get("tipo"), {"product", "service"}, "product")
        product.price = _decimal(row.get("precio_general"))
        product.price_wholesale = _decimal(row.get("precio_mayorista"))
        product.price_special = _decimal(row.get("precio_especial"))
        product.cost = _decimal(row.get("costo"))
        product.stock = int(_decimal(row.get("stock")))
        product.track_stock = _bool(row.get("controla_stock"), True)
        product.track_batches = _bool(row.get("controla_lotes"), product.kind == "product")
        product.is_active = _bool(row.get("activo"), True)
        product.image_url = _text(row.get("imagen_url")) or None

        category_name = _text(row.get("categoria"))
        product.category = _category_for_name(tenant, category_name, category_by_name) if category_name else None
        stats["products_created" if created else "products_updated"] += 1


def _import_customers(ws, tenant, stats) -> None:
    for rownum, row in _rows(ws, CUSTOMER_HEADERS, "Clientes", stats):
        name = _text(row.get("nombre"))
        if not name:
            _error(stats, f"Clientes fila {rownum}: nombre es obligatorio.")
            continue

        tax_id = _text(row.get("rtn"))
        customer = Customer.query.filter_by(tenant_id=tenant.id, tax_id=tax_id).first() if tax_id else None
        created = customer is None
        if created:
            customer = Customer(tenant_id=tenant.id)
            db.session.add(customer)

        customer.name = name
        customer.tax_id = tax_id or None
        customer.email = _text(row.get("email")) or None
        customer.phone = _text(row.get("telefono")) or None
        customer.address = _text(row.get("direccion")) or None
        customer.city = _text(row.get("ciudad")) or None
        customer.preferred_price_tier = _choice(row.get("tarifa"), {"general", "mayorista", "especial"}, "general")
        customer.country_code = (_text(row.get("pais")) or tenant.country_code or "HN")[:2].upper()
        customer.notes = _text(row.get("notas")) or None
        customer.is_active = _bool(row.get("activo"), True)
        stats["customers_created" if created else "customers_updated"] += 1


def _import_suppliers(ws, tenant, stats) -> None:
    for rownum, row in _rows(ws, SUPPLIER_HEADERS, "Proveedores", stats):
        name = _text(row.get("nombre"))
        if not name:
            _error(stats, f"Proveedores fila {rownum}: nombre es obligatorio.")
            continue

        code = _text(row.get("codigo"))
        tax_id = _text(row.get("rtn"))
        supplier = Supplier.query.filter_by(tenant_id=tenant.id, code=code).first() if code else None
        if supplier is None and tax_id:
            supplier = Supplier.query.filter_by(tenant_id=tenant.id, tax_id=tax_id).first()
        created = supplier is None
        if created:
            supplier = Supplier(tenant_id=tenant.id, code=code or _next_supplier_code(tenant.id))
            db.session.add(supplier)

        supplier.name = name
        supplier.legal_name = _text(row.get("razon_social")) or None
        supplier.tax_id = tax_id or None
        supplier.contact_name = _text(row.get("contacto")) or None
        supplier.email = _text(row.get("email")) or None
        supplier.phone = _text(row.get("telefono")) or None
        supplier.address = _text(row.get("direccion")) or None
        supplier.city = _text(row.get("ciudad")) or None
        supplier.country_code = (_text(row.get("pais")) or tenant.country_code or "HN")[:2].upper()
        supplier.payment_terms_days = int(_decimal(row.get("dias_credito")))
        supplier.credit_limit = _decimal(row.get("limite_credito"))
        supplier.price_includes_tax = _bool(row.get("precios_incluyen_impuesto"), False)
        supplier.notes = _text(row.get("notas")) or None
        supplier.is_active = _bool(row.get("activo"), True)
        stats["suppliers_created" if created else "suppliers_updated"] += 1


def _setup_sheet(ws, title: str, headers: list[str]) -> None:
    ws.title = title
    ws.append(headers)
    fill = PatternFill("solid", fgColor="EEF2FF")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1E293B")
        cell.fill = fill
    ws.freeze_panes = "A2"


def _rows(ws, headers, sheet_name, stats):
    header_map = {
        _text(cell.value).lower(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if _text(cell.value)
    }
    missing = [h for h in headers if h not in header_map]
    if missing:
        _error(stats, f"{sheet_name}: faltan columnas: {', '.join(missing)}.")
        return

    for rownum in range(2, ws.max_row + 1):
        values = {h: ws.cell(row=rownum, column=header_map[h]).value for h in headers}
        if any(_text(v) for v in values.values()):
            yield rownum, values


def _category_for_name(tenant, name: str, category_by_name):
    key = name.strip().lower()
    cat = category_by_name.get(key)
    if cat is None:
        cat = Category(tenant_id=tenant.id, name=name.strip())
        db.session.add(cat)
        category_by_name[key] = cat
    return cat


def _next_supplier_code(tenant_id: int) -> str:
    n = (Supplier.query.filter_by(tenant_id=tenant_id).count() or 0) + 1
    return f"P-{n:04d}"


def _to_stream(wb: Workbook) -> BytesIO:
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _autosize(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 42)


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal(value) -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else 0))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _bool(value, default: bool) -> bool:
    if value is None or value == "":
        return default
    return str(value).strip().lower() in {"1", "si", "sí", "s", "true", "activo", "yes", "y"}


def _choice(value, allowed: set[str], default: str) -> str:
    clean = _text(value).lower()
    return clean if clean in allowed else default


def _yes_no(value) -> str:
    return "si" if bool(value) else "no"


def _error(stats, message: str) -> None:
    if len(stats["errors"]) < 20:
        stats["errors"].append(message)
